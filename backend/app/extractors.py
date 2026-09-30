import io
from dataclasses import dataclass, field

import docx
import pymupdf  # PyMuPDF
from docx.oxml.ns import qn


@dataclass
class Extracted:
    text: str = ""
    meta: dict = field(default_factory=dict)
    images: list[tuple[bytes, str]] = field(default_factory=list)  # (bytes, mime)


def extract_pdf(data: bytes) -> Extracted:
    doc = pymupdf.open(stream=data, filetype="pdf")
    text = "\n".join(page.get_text("text") for page in doc).strip()

    tables = 0
    for page in doc:
        try:
            tables += len(page.find_tables().tables)
        except Exception:
            pass

    meta = {
        "pages": len(doc),
        "embedded_images": sum(len(p.get_images()) for p in doc),
        "tables_detected": tables,
        "fonts": sorted({f[3] for p in doc for f in p.get_fonts()}),
        "text_extractable": len(text) > 100,
    }

    images: list[tuple[bytes, str]] = []
    if not meta["text_extractable"]:  # scanned PDF -> send pages as images to Gemini
        for page in list(doc)[:3]:
            images.append((page.get_pixmap(dpi=150).tobytes("png"), "image/png"))

    return Extracted(text=text, meta=meta, images=images)


def extract_docx(data: bytes) -> Extracted:
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]

    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))

    columns = 1
    for s in d.sections:
        cols = s._sectPr.find(qn("w:cols"))
        if cols is not None and cols.get(qn("w:num")):
            columns = max(columns, int(cols.get(qn("w:num"))))

    header_footer_text = any(
        p.text.strip()
        for s in d.sections
        for p in list(s.header.paragraphs) + list(s.footer.paragraphs)
    )

    meta = {
        "tables_detected": len(d.tables),
        "column_count": columns,
        "embedded_images": len(d.inline_shapes),
        "text_in_header_footer": header_footer_text,
        "text_extractable": True,
    }
    return Extracted(text="\n".join(parts), meta=meta)


def extract_image(data: bytes, mime: str) -> Extracted:
    return Extracted(
        meta={"input_type": "image", "text_extractable": False},
        images=[(data, mime)],
    )
